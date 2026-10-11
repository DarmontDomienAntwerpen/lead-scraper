"""Export node: Pydantic validation + Excel export.

Applies the suppression list (compliance.py), builds the Pydantic `Lead` model for
every approved lead, and writes an Excel file with two sheets: "Leads" (for manual
review and sending) and "audit_log" (full traceability: every lead, including
rejected or failed ones, with source and timestamp).
"""

from __future__ import annotations

import logging
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from src.compliance import apply_suppression, load_suppression_domains
from src.config import OUTPUT_DIR
from src.state import LeadState

logger = logging.getLogger(__name__)

_COLUMN_WIDTHS = {
    "Company": 28,
    "Website": 34,
    "Phone": 16,
    "Email": 30,
    "Email status": 16,
    "Contact form": 34,
    "Niche": 24,
    "Opening line": 50,
    "Personalised email": 70,
    "Subject line": 30,
    "Need score": 14,
    "Hook": 50,
    "Signals": 60,
}

_LOG_HEADERS = ["Company", "Search term", "Source", "Fetched at", "Status", "Reason"]


class Lead(BaseModel):
    """One approved lead: exactly the columns, in order, that the user reviews manually."""

    model_config = ConfigDict(populate_by_name=True)

    company_name: str = Field(alias="Company")
    website: str = Field(alias="Website")
    phone: str = Field(default="", alias="Phone")
    email: str = Field(default="", alias="Email")
    email_status: str = Field(default="", alias="Email status")
    contact_form: str = Field(default="", alias="Contact form")
    niche: str = Field(alias="Niche")
    opening_line: str = Field(alias="Opening line")
    mail: str = Field(alias="Personalised email")
    subject_line: str = Field(default="", alias="Subject line")
    need_score: int = Field(default=0, alias="Need score")
    hook: str = Field(default="", alias="Hook")
    signal_flags: str = Field(default="", alias="Signals")

    @classmethod
    def from_lead_state(cls, lead: LeadState) -> "Lead":
        active_signals = ", ".join(
            signal_name for signal_name, present in (lead.get("signal_flags") or {}).items() if present
        )
        return cls(
            company_name=lead["company_name"],
            website=lead["website"] or "",
            phone=lead["phone"] or "",
            email=lead["email"] or "",
            email_status=lead["email_status"] or "",
            contact_form=lead["contact_url"] or "",
            niche=lead["niche"],
            opening_line=lead["icebreaker"],
            mail=lead["mail"],
            subject_line=lead.get("subject_line") or "",
            need_score=lead.get("need_score") or 0,
            hook=lead.get("hook") or "",
            signal_flags=active_signals,
        )


COLUMN_HEADERS: list[str] = [field.alias for field in Lead.model_fields.values()]


def export_leads(leads: list[LeadState], region: str, output_dir: str = OUTPUT_DIR) -> str:
    """Apply compliance rules, write the Excel file and return its path."""
    suppression_domains = load_suppression_domains()
    apply_suppression(leads, suppression_domains)

    approved = [Lead.from_lead_state(lead) for lead in leads if lead["status"] == "ok"]

    workbook = Workbook()
    leads_sheet = workbook.active
    leads_sheet.title = "Leads"
    _write_leads_sheet(leads_sheet, approved)

    log_sheet = workbook.create_sheet("audit_log")
    _write_log_sheet(log_sheet, leads)

    os.makedirs(output_dir, exist_ok=True)
    filename = f"leads_{region.lower()}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    path = os.path.join(output_dir, filename)
    workbook.save(path)

    logger.info(
        "Excel exported to %s (%d approved leads out of %d)",
        path,
        len(approved),
        len(leads),
    )
    return path


def _write_leads_sheet(sheet, approved_leads: list[Lead]) -> None:
    sheet.append(COLUMN_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for idx, header in enumerate(COLUMN_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = _COLUMN_WIDTHS.get(header, 20)

    for lead in approved_leads:
        values = lead.model_dump(by_alias=True)
        sheet.append([values[header] for header in COLUMN_HEADERS])

    mail_col = COLUMN_HEADERS.index("Personalised email") + 1
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=mail_col, max_col=mail_col):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"


def _write_log_sheet(sheet, all_leads: list[LeadState]) -> None:
    sheet.append(_LOG_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for idx, width in enumerate([28, 18, 16, 24, 12, 45], start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    for lead in all_leads:
        sheet.append(
            [
                lead["company_name"],
                lead["search_term"],
                lead["source"],
                lead["fetched_at"],
                lead["status"],
                lead.get("error_reason") or "",
            ]
        )

    sheet.freeze_panes = "A2"
