from openpyxl import load_workbook

from src.nodes.export import COLUMN_HEADERS, export_leads
from src.state import new_lead


def _make_lead(status: str = "ok", website: str = "https://peeters-fietsen.be") -> dict:
    lead = new_lead(
        search_term="fietsenmaker",
        place_id="p1",
        company_name="Fietsen Peeters",
        website=website,
        phone="03 111 22 33",
        source="Google Places",
        fetched_at="2026-06-18T10:00:00",
    )
    lead["niche"] = "fietsenmaker gespecialiseerd in elektrische fietsen"
    lead["icebreaker"] = "Ik zag dat jullie e-bikes herstellen."
    lead["mail"] = "Korte mail tekst."
    lead["subject_line"] = "Offertes overtypen?"
    lead["need_score"] = 3
    lead["hook"] = "Intake loopt via e-mail."
    lead["signal_flags"] = {"mailto_contact": True, "has_form": False}
    lead["status"] = status
    return lead


def test_export_leads_writes_expected_columns_and_filters_status(tmp_path):
    ok_lead = _make_lead(status="ok")
    skipped_lead = _make_lead(status="skipped", website="https://andere.be")
    skipped_lead["place_id"] = "p2"

    path = export_leads([ok_lead, skipped_lead], region="Antwerpen", output_dir=str(tmp_path))

    wb = load_workbook(path)
    assert wb.sheetnames == ["Leads", "audit_log"]

    leads_sheet = wb["Leads"]
    header = [cell.value for cell in leads_sheet[1]]
    assert header == COLUMN_HEADERS == [
        "Company",
        "Website",
        "Phone",
        "Email",
        "Email status",
        "Contact form",
        "Niche",
        "Opening line",
        "Personalised email",
        "Subject line",
        "Need score",
        "Hook",
        "Signals",
    ]
    assert leads_sheet.max_row == 2  # header + 1 ok-lead
    assert leads_sheet.cell(row=2, column=1).value == "Fietsen Peeters"
    assert leads_sheet.freeze_panes == "A2"

    row = {h: c.value for h, c in zip(header, leads_sheet[2])}
    assert row["Subject line"] == "Offertes overtypen?"
    assert row["Need score"] == 3
    assert row["Hook"] == "Intake loopt via e-mail."
    assert row["Signals"] == "mailto_contact"

    log_sheet = wb["audit_log"]
    assert log_sheet.max_row == 3  # header + 2 leads (ok + skipped)


def test_export_leads_applies_suppression(tmp_path):
    suppression_path = tmp_path / "suppression.csv"
    suppression_path.write_text("domein\npeeters-fietsen.be\n")

    import src.nodes.export as export_module

    original = export_module.load_suppression_domains
    export_module.load_suppression_domains = lambda path=None: original(str(suppression_path))
    try:
        lead = _make_lead(status="ok")
        path = export_leads([lead], region="Antwerpen", output_dir=str(tmp_path))
    finally:
        export_module.load_suppression_domains = original

    wb = load_workbook(path)
    assert wb["Leads"].max_row == 1  # enkel header, lead werd onderdrukt


def test_export_leads_includes_contact_url_when_no_email(tmp_path):
    lead = _make_lead(status="ok")
    lead["email"] = None
    lead["contact_url"] = "https://peeters-fietsen.be/contact"

    path = export_leads([lead], region="Antwerpen", output_dir=str(tmp_path))

    wb = load_workbook(path)
    sheet = wb["Leads"]
    headers = [cell.value for cell in sheet[1]]
    row = {h: c for h, c in zip(headers, [cell.value for cell in sheet[2]])}
    assert not row["Email"]
    assert row["Contact form"] == "https://peeters-fietsen.be/contact"


def test_export_leads_includes_email_status(tmp_path):
    lead = _make_lead(status="ok")
    lead["email"] = "info@peeters-fietsen.be"
    lead["email_status"] = "catchall"

    path = export_leads([lead], region="Antwerpen", output_dir=str(tmp_path))

    wb = load_workbook(path)
    sheet = wb["Leads"]
    headers = [cell.value for cell in sheet[1]]
    row = {h: c for h, c in zip(headers, [cell.value for cell in sheet[2]])}
    assert row["Email"] == "info@peeters-fietsen.be"
    assert row["Email status"] == "catchall"
