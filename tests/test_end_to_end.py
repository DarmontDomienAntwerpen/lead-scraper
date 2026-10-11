"""End-to-end test of the full StateGraph with every external call mocked."""

from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook

import src.graph as graph_module
from src.nodes.export import export_leads as real_export_leads
from src.state import new_lead


def _export_to_tmp(tmp_path):
    def _export(leads, region):
        return real_export_leads(leads, region, output_dir=str(tmp_path))

    return _export


def _stub_lead(place_id: str, with_website: bool = True) -> dict:
    return new_lead(
        search_term="loodgieter",
        place_id=place_id,
        company_name=f"Bedrijf {place_id}",
        website="https://example.be" if with_website else None,
        phone="03 000 00 00",
        source="Google Places",
        fetched_at="2026-06-18T10:00:00",
    )


@pytest.mark.asyncio
async def test_full_graph_runs_search_to_export_with_mixed_outcomes(tmp_path):
    leads = [_stub_lead("p1"), _stub_lead("p2"), _stub_lead("p3")]

    async def fake_search(search_term, region, max_leads):
        return leads

    async def fake_scrape(batch):
        # A phone number without a booking tool gives phone_only_cta, which reaches the
        # default need-score threshold (the signals node really runs in this test).
        batch[0]["raw_text"] = "Echte inhoud over loodgieterswerk. Bel ons op 03 111 22 33."
        batch[0]["raw_html"] = batch[0]["raw_text"]
        batch[1]["status"] = "error"
        batch[1]["error_reason"] = "Timeout"
        batch[2]["raw_text"] = "Cookie-muur pagina."
        return batch

    async def fake_analyze(batch):
        for lead in batch:
            if lead["status"] != "pending":
                continue
            if "Cookie" in lead["raw_text"]:
                lead["status"] = "skipped"
                lead["error_reason"] = "Bad data"
            else:
                lead["niche"] = "loodgieter gespecialiseerd in lekdetectie"
                lead["icebreaker"] = "Ik zag dat jullie lekdetectie aanbieden."
                lead["status"] = "ok"
        return batch

    async def fake_write(batch):
        for lead in batch:
            if lead["status"] == "ok":
                lead["mail"] = (
                    "Ik zag dat jullie lekdetectie aanbieden. "
                    "Met automatisering hou je veel tijd over op planning. "
                    "Mag ik een kort gesprek inplannen om te horen wat bij jullie veel tijd kost? "
                    "Laat het gerust weten als je dit niet wil ontvangen."
                )
        return batch

    with patch.object(graph_module, "search_leads", AsyncMock(side_effect=fake_search)), \
         patch.object(graph_module, "scrape_leads", AsyncMock(side_effect=fake_scrape)), \
         patch.object(graph_module, "analyze_leads", AsyncMock(side_effect=fake_analyze)), \
         patch.object(graph_module, "write_mails", AsyncMock(side_effect=fake_write)), \
         patch.object(graph_module, "export_leads", side_effect=_export_to_tmp(tmp_path)):
        compiled = graph_module.build_graph()
        final_state = await compiled.ainvoke(
            {"search_term": "loodgieter", "region": "Antwerpen", "max_leads": 20, "leads": [], "output_path": None}
        )

    assert final_state["output_path"]
    wb = load_workbook(final_state["output_path"])
    assert wb["Leads"].max_row == 2  # header + de ene geslaagde lead (p1)
    assert wb["Leads"].cell(row=2, column=1).value == "Bedrijf p1"
    assert wb["audit_log"].max_row == 4  # header + alle 3 leads


@pytest.mark.asyncio
async def test_full_graph_skips_to_export_when_search_finds_nothing(tmp_path):
    async def fake_search(search_term, region, max_leads):
        return []

    with patch.object(graph_module, "search_leads", AsyncMock(side_effect=fake_search)), \
         patch.object(graph_module, "export_leads", side_effect=_export_to_tmp(tmp_path)):
        compiled = graph_module.build_graph()
        final_state = await compiled.ainvoke(
            {"search_term": "x", "region": "Antwerpen", "max_leads": 20, "leads": [], "output_path": None}
        )

    wb = load_workbook(final_state["output_path"])
    assert wb["Leads"].max_row == 1  # enkel header
